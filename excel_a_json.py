import openpyxl
import json
import os

def convertir_excel_a_json():
    excel_path = r"C:\Users\seban\AndroidStudioProjects\Cannon\Plantilla_Catalogo_Cannon.xlsx"
    json_path = r"C:\Users\seban\AndroidStudioProjects\Cannon\precios_azure.json"
    
    if not os.path.exists(excel_path):
        print(f"Error: No se encontró la planilla en {excel_path}")
        return False
        
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    # Usar la hoja 'Resumen General' o la primera hoja
    ws = wb["Resumen General"] if "Resumen General" in wb.sheetnames else wb.active
    
    # Estructura: nombre -> { "nombre": str, "descuento": float, "plazas": { plaza: precio } }
    productos_dict = {}
    
    # Filas: Categoría (1), Producto (2), Plaza (3), Precio Normal (4), Descuento (5)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
            
        nombre = str(row[1]).strip()
        plaza = str(row[2]).strip() if row[2] else "Única"
        
        try:
            precio = float(row[3])
        except (ValueError, TypeError):
            continue
            
        try:
            dcto_val = float(row[4])
            # Si en Excel se guardó como 0.50 (50%), multiplicar por 100
            if dcto_val <= 1.0 and dcto_val > 0:
                dcto = dcto_val * 100.0
            else:
                dcto = dcto_val
        except (ValueError, TypeError):
            dcto = 0.0
            
        if nombre not in productos_dict:
            productos_dict[nombre] = {
                "nombre": nombre,
                "descuento": round(dcto, 1),
                "plazas": {}
            }
            
        productos_dict[nombre]["plazas"][plaza] = precio
        
    lista_final = list(productos_dict.values())
    
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(lista_final, f, ensure_ascii=False, indent=2)
        
    print(f"✅ ¡Éxito! Se procesaron {len(lista_final)} productos desde Excel a {json_path}")
    return True

if __name__ == "__main__":
    convertir_excel_a_json()
