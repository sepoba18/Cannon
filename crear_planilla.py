import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import shutil
import os

with open('precios_azure.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

def get_cat(nombre):
    n = nombre.upper()
    if 'QUILT' in n or 'Q.' in n:
        return 'Quilts'
    if 'SAB' in n or 'VIVO' in n or 'TENCEL' in n:
        return 'Sábanas'
    if 'PLUMON' in n:
        return 'Plumones'
    if 'ALMOHADA' in n or 'COPPER' in n:
        return 'Almohadas'
    if 'TOPPER' in n or 'IMPERMEABLE' in n:
        return 'Toppers'
    if 'FDA' in n or 'FUNDA' in n:
        return 'Fundas'
    return 'Otros'

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Catálogo Cannon'
ws.views.sheetView[0].showGridLines = True

headers = ['Producto', 'Categoría', 'Plaza', 'Precio Normal', 'Descuento (%)']
ws.append(headers)

header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='0B192C', end_color='0B192C', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center')

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
ws.row_dimensions[1].height = 28

items_sorted = sorted(items, key=lambda x: (get_cat(x['nombre']), x['nombre']))
plaza_order = {'1.0': 1, '1.5': 2, '2.0': 3, '2.5': 4, '3.0': 5, 'King': 6, 'Super King': 7, 'Única': 8, 'MAISON': 9}

row_idx = 2
for it in items_sorted:
    cat = get_cat(it['nombre'])
    sorted_plazas = sorted(it['plazas'].items(), key=lambda p: plaza_order.get(p[0], 99))
    for plaza, precio in sorted_plazas:
        c_prod = ws.cell(row=row_idx, column=1, value=it['nombre'])
        c_cat = ws.cell(row=row_idx, column=2, value=cat)
        c_plaza = ws.cell(row=row_idx, column=3, value=plaza)
        c_precio = ws.cell(row=row_idx, column=4, value=int(precio))
        c_dcto = ws.cell(row=row_idx, column=5, value=int(it['descuento']))

        c_prod.font = Font(name='Segoe UI', size=10, bold=True)
        c_cat.font = Font(name='Segoe UI', size=10)
        c_plaza.font = Font(name='Segoe UI', size=10)
        c_precio.font = Font(name='Segoe UI', size=10)
        c_dcto.font = Font(name='Segoe UI', size=10, bold=True, color='007A33')

        c_cat.alignment = Alignment(horizontal='center')
        c_plaza.alignment = Alignment(horizontal='center')
        c_precio.alignment = Alignment(horizontal='right')
        c_dcto.alignment = Alignment(horizontal='center')

        c_precio.number_format = '$ #,##0'
        c_dcto.number_format = '0"%"'

        for col in range(1, 6):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

        ws.row_dimensions[row_idx].height = 20
        row_idx += 1

ws.column_dimensions['A'].width = 34
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 18
ws.column_dimensions['E'].width = 18

ws.auto_filter.ref = f'A1:E{row_idx - 1}'

output_path = 'Plantilla_Catalogo_Cannon_12.xlsx'
wb.save(output_path)
print(f'Guardado localmente en: {output_path} con {row_idx - 1} filas')

gdrive_path = r'G:\Mi unidad\Antigravity_Docs\Plantilla_Catalogo_Cannon_12.xlsx'
if os.path.exists(r'G:\Mi unidad\Antigravity_Docs'):
    shutil.copyfile(output_path, gdrive_path)
    print(f'Copia sincronizada a Google Drive: {gdrive_path}')
